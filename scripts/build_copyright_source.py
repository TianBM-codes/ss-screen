"""Build deterministic source-code identification material for registration."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

SOFTWARE_NAME = "新材料计算筛选软件"
SOFTWARE_VERSION = "V0.1.0"
DEFAULT_LINES_PER_SECTION = 2500
LINES_PER_PAGE = 50
MONO_FONT_PATH = Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")
BOLD_FONT_PATH = Path("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf")
LATIN_BOLD_FONT_PATH = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
MONO_FONT = "SourceMono"
BOLD_FONT = "SourceHeader"
LATIN_BOLD_FONT = "SourceLatinHeader"


@dataclass(frozen=True)
class SourceLine:
    sequence: int
    path: str
    source_line: int
    text: str


def collect_source_lines(source_root: Path) -> tuple[list[Path], list[SourceLine]]:
    files = sorted(source_root.rglob("*.py"), key=lambda path: path.as_posix())
    lines: list[SourceLine] = []
    for path in files:
        relative_path = path.as_posix()
        for source_line, text in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
            lines.append(
                SourceLine(
                    sequence=len(lines) + 1,
                    path=relative_path,
                    source_line=source_line,
                    text=text,
                )
            )
    return files, lines


def select_lines(lines: list[SourceLine], count: int) -> list[tuple[str, SourceLine]]:
    if count <= 0:
        raise ValueError("line count must be positive")
    if len(lines) < count * 2:
        raise ValueError(f"source has {len(lines)} lines; at least {count * 2} are required")
    return [("FRONT", line) for line in lines[:count]] + [("BACK", line) for line in lines[-count:]]


def display_line(section: str, section_line: int, line: SourceLine) -> str:
    return (
        f"{section[0]}{section_line:04d} "
        f"{line.path}:{line.source_line:04d} | {line.text.expandtabs(4)}"
    )


def display_pdf_line(section: str, section_line: int, line: SourceLine) -> str:
    return f"{section[0]}{section_line:04d} {line.source_line:04d} | {line.text.expandtabs(4)}"


def write_text(output: Path, selected: list[tuple[str, SourceLine]]) -> None:
    counters = {"FRONT": 0, "BACK": 0}
    rendered: list[str] = []
    for section, line in selected:
        counters[section] += 1
        rendered.append(display_line(section, counters[section], line))
    output.write_text("\n".join(rendered) + "\n", encoding="utf-8")


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, field, separate, text, end):
        run._r.append(element)


def _configure_docx_section(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(12)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{SOFTWARE_NAME} {SOFTWARE_VERSION}  源程序  第")
    run.font.name = "宋体"
    run.font.size = Pt(8)
    _add_field(paragraph, "PAGE")
    run = paragraph.add_run("页/共")
    run.font.name = "宋体"
    run.font.size = Pt(8)
    _add_field(paragraph, "NUMPAGES")
    run = paragraph.add_run("页")
    run.font.name = "宋体"
    run.font.size = Pt(8)


def _set_code_run_font(run, text: str) -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(6.5)
    run_properties = run._r.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Consolas")
    fonts.set(qn("w:hAnsi"), "Consolas")
    fonts.set(qn("w:eastAsia"), "宋体")
    visible_length = max(1, len(text.expandtabs(4)))
    if visible_length > 150:
        scale = max(35, round(150 / visible_length * 100))
        width = OxmlElement("w:w")
        width.set(qn("w:val"), str(scale))
        run_properties.append(width)


def _normalize_docx_archive(path: Path) -> None:
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with (
            zipfile.ZipFile(path, "r") as source,
            zipfile.ZipFile(
                temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
            ) as target,
        ):
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                target.writestr(info, source.read(name))
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def write_docx(output: Path, selected: list[tuple[str, SourceLine]]) -> None:
    document = Document()
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)
    document.core_properties.title = f"{SOFTWARE_NAME} {SOFTWARE_VERSION} 源程序"
    document.core_properties.subject = "源程序前后各2500行"
    document.core_properties.author = "SS-Screen"
    fixed_time = datetime(2000, 1, 1, tzinfo=UTC)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time
    section = document.sections[0]
    _configure_docx_section(section)

    for index, (_, line) in enumerate(selected):
        paragraph = document.add_paragraph()
        if index and index % LINES_PER_PAGE == 0:
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(13.7)
        run = paragraph.add_run(line.text.expandtabs(4))
        _set_code_run_font(run, line.text)

    document.save(output)
    _normalize_docx_archive(output)


def write_pdf(output: Path, selected: list[tuple[str, SourceLine]]) -> None:
    if not all(path.is_file() for path in (MONO_FONT_PATH, BOLD_FONT_PATH, LATIN_BOLD_FONT_PATH)):
        raise FileNotFoundError("required Unicode fonts are unavailable")
    pdfmetrics.registerFont(TTFont(MONO_FONT, MONO_FONT_PATH))
    pdfmetrics.registerFont(TTFont(BOLD_FONT, BOLD_FONT_PATH))
    pdfmetrics.registerFont(TTFont(LATIN_BOLD_FONT, LATIN_BOLD_FONT_PATH))
    page_width, page_height = A4
    left = 24
    right = 24
    top = 31
    bottom = 24
    header_height = 27
    font_size = 7.1
    leading = (page_height - top - bottom - header_height) / LINES_PER_PAGE
    pdf = canvas.Canvas(str(output), pagesize=A4, pageCompression=1, invariant=1)
    pdf.setTitle(f"{SOFTWARE_NAME} {SOFTWARE_VERSION} Source Code")
    total_pages = len(selected) // LINES_PER_PAGE
    counters = {"FRONT": 0, "BACK": 0}

    for offset in range(0, len(selected), LINES_PER_PAGE):
        page_lines = selected[offset : offset + LINES_PER_PAGE]
        page_number = offset // LINES_PER_PAGE + 1
        section = page_lines[0][0]
        pdf.setFont(BOLD_FONT, 8)
        pdf.drawString(left, page_height - top, SOFTWARE_NAME)
        title_width = stringWidth(SOFTWARE_NAME, BOLD_FONT, 8)
        pdf.setFont(LATIN_BOLD_FONT, 8)
        pdf.drawString(left + title_width + 6, page_height - top, SOFTWARE_VERSION)
        pdf.drawRightString(
            page_width - right,
            page_height - top,
            f"{section} SOURCE  Page {page_number}/{total_pages}",
        )
        pdf.line(left, page_height - top - 4, page_width - right, page_height - top - 4)
        first_path = page_lines[0][1].path
        last_path = page_lines[-1][1].path
        path_label = first_path if first_path == last_path else f"{first_path} -> {last_path}"
        pdf.setFont(MONO_FONT, 6.3)
        pdf.drawString(left, page_height - top - 13, path_label)
        pdf.setFont(MONO_FONT, font_size)
        y = page_height - top - header_height
        for item_section, line in page_lines:
            counters[item_section] += 1
            rendered = display_pdf_line(item_section, counters[item_section], line)
            rendered_width = stringWidth(rendered, MONO_FONT, font_size)
            available_width = page_width - left - right
            horizontal_scale = min(1.0, available_width / rendered_width)
            pdf.saveState()
            pdf.translate(left, y)
            pdf.scale(horizontal_scale, 1.0)
            pdf.drawString(0, 0, rendered)
            pdf.restoreState()
            y -= leading
        pdf.showPage()
    pdf.save()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build(source_root: Path, output_dir: Path, count: int) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    files, lines = collect_source_lines(source_root)
    selected = select_lines(lines, count)
    stem = "SS-Screen_V0.1.0_source_front2500_back2500"
    text_path = output_dir / f"{stem}.txt"
    pdf_path = output_dir / f"{stem}.pdf"
    docx_path = output_dir / f"{stem}_unlabeled.docx"
    manifest_path = output_dir / f"{stem}_manifest.json"
    write_text(text_path, selected)
    write_pdf(pdf_path, selected)
    write_docx(docx_path, selected)
    manifest = {
        "software_name": SOFTWARE_NAME,
        "software_version": SOFTWARE_VERSION,
        "source_root": source_root.as_posix(),
        "ordering": "POSIX path ascending, then original line number ascending",
        "source_file_count": len(files),
        "source_line_count": len(lines),
        "selected_front_lines": count,
        "selected_back_lines": count,
        "overlap": False,
        "lines_per_pdf_page": LINES_PER_PAGE,
        "pdf_page_count": len(selected) // LINES_PER_PAGE,
        "docx_body_line_count": len(selected),
        "docx_lines_per_page": LINES_PER_PAGE,
        "docx_line_prefixes": False,
        "first_selected_source": f"{selected[0][1].path}:{selected[0][1].source_line}",
        "last_front_source": f"{selected[count - 1][1].path}:{selected[count - 1][1].source_line}",
        "first_back_source": f"{selected[count][1].path}:{selected[count][1].source_line}",
        "last_selected_source": f"{selected[-1][1].path}:{selected[-1][1].source_line}",
        "source_files": [path.as_posix() for path in files],
        "outputs": {
            text_path.name: sha256(text_path),
            pdf_path.name: sha256(pdf_path),
            docx_path.name: sha256(docx_path),
        },
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=Path("src/ssscreen"))
    parser.add_argument("--output-dir", type=Path, default=Path("docs/copyright-source"))
    parser.add_argument("--lines-per-section", type=int, default=DEFAULT_LINES_PER_SECTION)
    args = parser.parse_args()
    build(args.source_root, args.output_dir, args.lines_per_section)


if __name__ == "__main__":
    main()
