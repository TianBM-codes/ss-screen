"""Update the main-function field in the copyright application DOCX."""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document

FIELD_PREFIX = "主要功能（500-1300字）："
MAIN_FUNCTIONS = (
    "本软件面向材料科学、半导体材料与计算材料学研究，用于从大规模材料数据中自动筛选具有结构相容性、带隙可调潜力和稳定性证据的固溶体或合金候选体系，并形成可追溯、可复核的标准化结果。"
    "一、软件可读取Materials Project在线接口或离线快照以及WBM等材料数据，统一整理材料编号、化学式、元素组成、晶体结构、带隙、形成能和凸包能等字段，记录数据来源、查询条件、数据库版本及文件校验信息；同时支持按元素数量、稳定性、带隙范围、原子数、价态可判定性和排除元素等可配置条件完成候选预筛。"
    "二、软件可生成或读取robocrys局部配位环境描述，将拟发生合金替换的可变元素统一映射为占位符X，比较配位数、近邻元素、连接方式和几何环境等结构指纹，把化学组成不同但局部结构原型相近的材料归入同一结构组，从而减少仅凭化学式或空间群判断造成的误配。"
    "三、软件能够为候选端元生成确定性任务编号和结构哈希，导出JSON、CIF、POSCAR及结果模板，供用户在VASP、ABACUS、AiiDA或其他外部平台开展HSE06、mBJ等高精度带隙计算；返回结果进入软件后，可校验任务编号、材料编号、计算方法、结构身份、参数设置、数值范围、重复记录、失败状态和覆盖率，避免不同结构或计算条件的结果被错误混用。"
    "四、软件可按近零带隙端、小带隙端和可调窗口等具名阈值回填带隙数据并枚举材料对，保留直接或间接带隙信息和数据来源，输出候选对、缺失记录、拒绝原因及方法间比较结果。"
    "五、对于进入稳定性预筛的材料对，软件支持随机替换或icet方法生成指定组分和超胞大小的特殊准随机结构，并同步导出两个端元结构、实际组分、随机种子和生成参数，保证后续计算能够复现。"
    "六、软件可调用经明确指定的MACE机器学习势，对端元、合金结构及竞争相执行可恢复的结构弛豫，记录模型名称与哈希、设备、能量、力、应力、收敛状态和失败原因；在同一模型与设置的能量基准下，计算合金相对于端元线性组合的混合焓，防止不同模型或不同参考态之间进行无效比较。"
    "七、软件结合Phonopy有限位移与机器学习势计算原子受力，生成声子频带、态密度及动力学稳定性摘要；还可从Materials Project在线接口或指定离线快照收集候选化学体系的竞争相结构，在统一机器学习势能量基准下构建凸包并计算候选结构的凸包距离，同时保留快照范围、结构缺失、弛豫失败和声子虚频等风险信息。"
    "八、软件按材料对和合金结构汇总高精度带隙、混合焓、声子、竞争相凸包及可选缺陷证据，检查跨阶段模型哈希与参数兼容性，给出promising、uncertain或low-priority三类研究优先级以及L2至L6证据等级，并列出缺失证据、警告和建议的下一步验证。"
    "各阶段均通过命令行参数接收输入输出路径和阈值，生成CSV、JSON、JSONL、结构文件与Markdown审计报告；任务标识、结构哈希、模型信息和失败语义贯穿流程，便于批量运行、断点续算、结果追踪、统计分析和人工复核。本软件用于科研候选预筛和决策支持，不以机器学习势结果替代最终第一性原理计算、实验验证或可合成性结论。"
)


def count_characters(text: str) -> int:
    """Count non-whitespace characters for the form's character limit."""
    return len(re.sub(r"\s+", "", text))


def replace_main_functions(input_path: Path, output_path: Path) -> int:
    document = Document(input_path)
    matched_cells = []
    seen_cells = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                identity = cell._tc
                if identity in seen_cells:
                    continue
                seen_cells.add(identity)
                if cell.text.strip().startswith("主要功能（500-1300字）"):
                    matched_cells.append(cell)

    if len(matched_cells) != 1:
        raise ValueError(f"expected one main-function cell, found {len(matched_cells)}")

    character_count = count_characters(MAIN_FUNCTIONS)
    if not 500 <= character_count <= 1300:
        raise ValueError(f"main-function text has {character_count} characters")

    cell = matched_cells[0]
    paragraph = cell.paragraphs[0]
    paragraph_properties = deepcopy(paragraph._p.pPr)
    run_properties = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_properties = deepcopy(run._r.rPr)
            break

    for child in list(cell._tc):
        if child.tag.endswith("}p"):
            cell._tc.remove(child)
    paragraph = cell.add_paragraph()
    if paragraph_properties is not None:
        paragraph._p.insert(0, paragraph_properties)
    run = paragraph.add_run(FIELD_PREFIX + MAIN_FUNCTIONS)
    if run_properties is not None:
        run._r.insert(0, run_properties)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return character_count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    count = replace_main_functions(args.input, args.output)
    print(f"wrote {args.output} with {count} main-function characters")


if __name__ == "__main__":
    main()
