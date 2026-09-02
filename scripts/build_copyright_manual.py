"""Build the SS-Screen software-copyright operation manual as DOCX and PDF."""

from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

from build_full_user_manual import (
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MID_GRAY,
    MUTED,
    NAVY,
    Block,
    add_body,
    add_page_field,
    configure_section,
    configure_styles,
    parse_markdown,
    set_cell_margins,
    set_paragraph_border,
    set_run_font,
)
from build_full_user_manual import (
    set_document_properties as _set_full_manual_properties,
)
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    XPreformatted,
)
from reportlab.platypus import Image as ReportLabImage
from reportlab.platypus.tableofcontents import TableOfContents

SOFTWARE_NAME = "新材料计算筛选软件"
SOFTWARE_SHORT_NAME = "SS-Screen"
SOFTWARE_VERSION = "V1.0"
MANUAL_TITLE = f"{SOFTWARE_NAME} {SOFTWARE_VERSION} 操作手册"
RIGHTS_HOLDER = "荚左龙、谢骁、任琦、朱博南"
AUTHOR = "SS-Screen Project"
BUILD_DATE = "2026年8月21日"
SUBMISSION_PLACEHOLDER = "申报主体名称（提交前替换）"

REQUIRED_HEADINGS = (
    "第一章 软件概述",
    "第二章 运行环境与安装",
    "第三章 项目目录与通用操作",
    "第四章 数据集构建与组成筛选",
    "第五章 结构描述归档",
    "第六章 结构匹配与材料分组",
    "第七章 高精度带隙任务交接",
    "第八章 端元配对与方法比较",
    "第九章 SQS合金结构生成",
    "第十章 MACE结构驰豫",
    "第十一章 混合焓计算",
    "第十二章 声子谱计算",
    "第十三章 竞争相与凸包分析",
    "第十四章 综合推荐",
    "第十五章 完整操作示例",
    "第十六章 结果解释与质量控制",
    "第十七章 常见问题与维护",
    "附录A 命令速查",
    "附录B 文件格式与字段",
    "附录C 文档版本记录",
)
FORBIDDEN_TEXT = (
    "完整版本预期接口",
    "当前代码尚未提供",
    "未来功能",
    "V1.4",
    "/home/zuolong/",
    "/vepfs-mlp2/",
)

PDF_FONT = "CopyrightCJK"
PDF_FONT_PATH = Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Markdown source")
    parser.add_argument("--docx", type=Path, required=True, help="Output DOCX")
    parser.add_argument("--pdf", type=Path, required=True, help="Output PDF")
    parser.add_argument(
        "--rights-holder",
        help="Replace the single rights-holder placeholder in generated outputs",
    )
    parser.add_argument("--check", action="store_true", help="Run structural checks")
    return parser.parse_args()


def validate_source(source: str) -> None:
    for heading in REQUIRED_HEADINGS:
        if f"# {heading}" not in source:
            raise ValueError(f"missing required heading: {heading}")
    for text in FORBIDDEN_TEXT:
        if text in source:
            raise ValueError(f"registration manual contains forbidden text: {text}")
    if source.count(SUBMISSION_PLACEHOLDER) > 1:
        raise ValueError("registration manual contains multiple rights-holder placeholders")
    if RIGHTS_HOLDER not in source and SUBMISSION_PLACEHOLDER not in source:
        raise ValueError("registration manual lacks the confirmed rights-holder names")
    if source.count(SOFTWARE_VERSION) < 2:
        raise ValueError("software version is not repeated in source identification fields")
    if source.count(SOFTWARE_NAME) < 3:
        raise ValueError("software full name is not consistently repeated")
    if "所有命令均来自 ss-screen 1.0" not in source:
        raise ValueError("manual lacks the frozen CLI-version declaration")
    if "0.1.0" in source:
        raise ValueError("V1.0 registration manual contains the obsolete version 0.1.0")


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    page_number = sect_pr.find(qn("w:pgNumType"))
    if page_number is None:
        page_number = OxmlElement("w:pgNumType")
        sect_pr.append(page_number)
    page_number.set(qn("w:start"), str(start))


def add_registration_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    set_page_number_start(section, 0)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = ""
    table = header.add_table(rows=1, cols=2, width=Cm(15.8))
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Cm(12.0)
    right.width = Cm(3.8)
    for cell in (left, right):
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    left_p = left.paragraphs[0]
    left_run = left_p.add_run(f"{SOFTWARE_NAME} {SOFTWARE_VERSION}")
    set_run_font(left_run, size=8.5, color=MUTED)
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(right_p)
    for run in right_p.runs:
        set_run_font(run, size=8.5, color=MUTED)
    set_paragraph_border(header_paragraph, color=MID_GRAY, size="4")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("软件著作权登记文档鉴别材料")
    set_run_font(footer_run, size=8.0, color=MUTED)

    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""


def add_registration_cover(doc: Document, rights_holder: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(22)
    run = kicker.add_run("软件著作权登记文档鉴别材料")
    set_run_font(run, size=12, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(15)
    run = title.add_run(SOFTWARE_NAME)
    set_run_font(run, size=27, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run(f"{SOFTWARE_VERSION} 操作手册")
    set_run_font(run, size=19, color=DARK_BLUE, bold=True)

    rule = doc.add_paragraph()
    set_paragraph_border(rule, color=BLUE, size="10")
    rule.paragraph_format.space_after = Pt(28)

    metadata = (
        ("软件简称", SOFTWARE_SHORT_NAME),
        ("软件版本", SOFTWARE_VERSION),
        ("软件类型", "材料科学计算与筛选命令行软件"),
        ("著作权人", rights_holder),
        ("编制日期", BUILD_DATE),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(3.0)
        paragraph.paragraph_format.right_indent = Cm(1.0)
        paragraph.paragraph_format.space_after = Pt(7)
        label_run = paragraph.add_run(f"{label}　")
        set_run_font(label_run, size=11, color=MUTED, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=11, color=INK)
    doc.add_page_break()


def set_registration_properties(doc: Document) -> None:
    _set_full_manual_properties(doc)
    properties = doc.core_properties
    properties.title = MANUAL_TITLE
    properties.subject = "SS-Screen 1.0 软件著作权登记操作手册"
    properties.author = AUTHOR
    properties.last_modified_by = AUTHOR
    properties.keywords = "SS-Screen, 软件著作权, 操作手册, 材料筛选"
    properties.comments = "Generated from the registration-manual Markdown source."


def build_docx(
    source: str,
    output: Path,
    *,
    rights_holder: str,
    source_dir: Path,
) -> None:
    blocks = parse_markdown(source)
    doc = Document()
    configure_section(doc)
    configure_styles(doc)
    set_registration_properties(doc)
    add_registration_header_footer(doc)
    add_registration_cover(doc, rights_holder)
    add_registration_toc(doc, blocks)
    add_body(doc, blocks, asset_base=source_dir)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = False
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def add_registration_toc(doc: Document, blocks: list[Block]) -> None:
    heading = doc.add_paragraph("目录")
    heading.style = doc.styles["Heading 1"]
    heading.paragraph_format.page_break_before = False
    used_lines = 2
    page_line_budget = 38
    for block in blocks:
        if block.kind != "heading" or block.level > 2:
            continue
        if block.level == 2 and re.match(r"^[BC]\.\d+", block.text):
            continue
        line_cost = 2 if block.level == 1 else 1
        if used_lines + line_cost > page_line_budget:
            doc.add_page_break()
            used_lines = 0
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0 if block.level == 1 else 0.7)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = block.level == 1
        run = paragraph.add_run(block.text)
        set_run_font(
            run,
            size=9.5 if block.level == 1 else 8.5,
            color=NAVY if block.level == 1 else INK,
            bold=block.level == 1,
        )
        used_lines += line_cost
    note = doc.add_paragraph("说明：目录按章节层级列出；页码以最终 PDF 为准。")
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in note.runs:
        set_run_font(run, size=8, color=MUTED, italic=True)
    doc.add_page_break()


def register_pdf_font() -> None:
    if not PDF_FONT_PATH.is_file():
        raise ValueError(f"required CJK font is missing: {PDF_FONT_PATH}")
    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT, PDF_FONT_PATH))
        pdfmetrics.registerFontFamily(
            PDF_FONT,
            normal=PDF_FONT,
            bold=PDF_FONT,
            italic=PDF_FONT,
            boldItalic=PDF_FONT,
        )


def inline_markup(text: str) -> str:
    pattern = re.compile(r"(\*\*.+?\*\*|\x60[^\x60]+\x60|\[[^\]]+\]\([^)]+\))")
    output: list[str] = []
    cursor = 0
    for match in pattern.finditer(text):
        output.append(html.escape(text[cursor : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            output.append(f"<b>{html.escape(token[2:-2])}</b>")
        elif token.startswith(chr(96)):
            output.append(
                f'<font name="{PDF_FONT}" color="#1F4D78">' f"{html.escape(token[1:-1])}</font>"
            )
        else:
            label = token[1 : token.index("]")]
            output.append(f'<font color="#2E74B5">{html.escape(label)}</font>')
        cursor = match.end()
    output.append(html.escape(text[cursor:]))
    return "".join(output)


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle(
            "CopyrightBody",
            parent=base["BodyText"],
            fontName=PDF_FONT,
            fontSize=9.25,
            leading=13,
            textColor=colors.HexColor(f"#{INK}"),
            spaceAfter=5,
            wordWrap="CJK",
        ),
        "h1": ParagraphStyle(
            "CopyrightHeading1",
            parent=base["Heading1"],
            fontName=PDF_FONT,
            fontSize=17,
            leading=24,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=6,
            spaceAfter=10,
            wordWrap="CJK",
        ),
        "h2": ParagraphStyle(
            "CopyrightHeading2",
            parent=base["Heading2"],
            fontName=PDF_FONT,
            fontSize=13,
            leading=19,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceBefore=10,
            spaceAfter=6,
            wordWrap="CJK",
        ),
        "h3": ParagraphStyle(
            "CopyrightHeading3",
            parent=base["Heading3"],
            fontName=PDF_FONT,
            fontSize=11,
            leading=17,
            textColor=colors.HexColor(f"#{DARK_BLUE}"),
            spaceBefore=8,
            spaceAfter=4,
            wordWrap="CJK",
        ),
        "code": ParagraphStyle(
            "CopyrightCode",
            parent=base["Code"],
            fontName=PDF_FONT,
            fontSize=7.2,
            leading=10,
            leftIndent=8,
            rightIndent=8,
            borderPadding=6,
            backColor=colors.HexColor(f"#{LIGHT_GRAY}"),
            textColor=colors.HexColor(f"#{INK}"),
            wordWrap="CJK",
            spaceBefore=4,
            spaceAfter=6,
        ),
        "toc_title": ParagraphStyle(
            "CopyrightTocTitle",
            parent=base["Title"],
            fontName=PDF_FONT,
            fontSize=20,
            leading=28,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=16,
        ),
        "cover_kicker": ParagraphStyle(
            "CopyrightCoverKicker",
            parent=base["BodyText"],
            fontName=PDF_FONT,
            fontSize=12,
            leading=18,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceAfter=20,
        ),
        "cover_title": ParagraphStyle(
            "CopyrightCoverTitle",
            parent=base["Title"],
            fontName=PDF_FONT,
            fontSize=26,
            leading=38,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=12,
            wordWrap="CJK",
        ),
        "cover_version": ParagraphStyle(
            "CopyrightCoverVersion",
            parent=base["Heading2"],
            fontName=PDF_FONT,
            fontSize=18,
            leading=26,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{DARK_BLUE}"),
            spaceAfter=26,
        ),
    }


class CopyrightPdfTemplate(BaseDocTemplate):
    def __init__(self, output: Path, styles: dict[str, ParagraphStyle]) -> None:
        super().__init__(
            str(output),
            pagesize=A4,
            leftMargin=2.35 * cm,
            rightMargin=2.35 * cm,
            topMargin=2.25 * cm,
            bottomMargin=2.25 * cm,
            title=MANUAL_TITLE,
            author=AUTHOR,
            subject="软件著作权登记操作手册",
        )
        self.manual_styles = styles
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="manual-frame",
        )
        self.addPageTemplates(PageTemplate(id="manual", frames=frame, onPage=self.draw_page))

    def draw_page(self, canvas, _doc) -> None:
        physical_page = canvas.getPageNumber()
        if physical_page == 1:
            return
        logical_page = physical_page - 1
        canvas.saveState()
        canvas.setFont(PDF_FONT, 8)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(
            self.leftMargin,
            A4[1] - 1.35 * cm,
            f"{SOFTWARE_NAME} {SOFTWARE_VERSION}",
        )
        canvas.drawRightString(
            A4[0] - self.rightMargin,
            A4[1] - 1.35 * cm,
            f"第 {logical_page} 页",
        )
        canvas.setStrokeColor(colors.HexColor(f"#{MID_GRAY}"))
        canvas.line(
            self.leftMargin,
            A4[1] - 1.52 * cm,
            A4[0] - self.rightMargin,
            A4[1] - 1.52 * cm,
        )
        canvas.drawCentredString(
            A4[0] / 2,
            1.25 * cm,
            "软件著作权登记文档鉴别材料",
        )
        canvas.restoreState()

    def afterFlowable(self, flowable) -> None:
        level = getattr(flowable, "_toc_level", None)
        if level is None:
            return
        text = flowable.getPlainText()
        key = f"heading-{self.seq.nextf('heading')}"
        self.canv.bookmarkPage(key)
        logical_page = max(1, self.page - 1)
        self.canv.addOutlineEntry(text, key, level=level, closed=False)
        self.notify("TOCEntry", (level, text, logical_page, key))


def cover_story(styles: dict[str, ParagraphStyle], *, rights_holder: str) -> list:
    metadata = [
        ["软件简称", SOFTWARE_SHORT_NAME],
        ["软件版本", SOFTWARE_VERSION],
        ["软件类型", "材料科学计算与筛选命令行软件"],
        ["著作权人", rights_holder],
        ["编制日期", BUILD_DATE],
    ]
    table = Table(metadata, colWidths=[3.2 * cm, 9.3 * cm], hAlign="CENTER")
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), PDF_FONT),
                ("FONTSIZE", (0, 0), (-1, -1), 10.5),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor(f"#{MUTED}")),
                ("TEXTCOLOR", (1, 0), (1, -1), colors.HexColor(f"#{INK}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor(f"#{MID_GRAY}")),
            ]
        )
    )
    return [
        Spacer(1, 3.4 * cm),
        Paragraph("软件著作权登记文档鉴别材料", styles["cover_kicker"]),
        Paragraph(SOFTWARE_NAME, styles["cover_title"]),
        Paragraph(f"{SOFTWARE_VERSION} 操作手册", styles["cover_version"]),
        table,
        PageBreak(),
    ]


def toc_story(styles: dict[str, ParagraphStyle]) -> list:
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle(
            "CopyrightTocLevel1",
            fontName=PDF_FONT,
            fontSize=9.5,
            leading=15,
            leftIndent=0,
            firstLineIndent=0,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=2,
        ),
        ParagraphStyle(
            "CopyrightTocLevel2",
            fontName=PDF_FONT,
            fontSize=7.8,
            leading=10.5,
            leftIndent=16,
            firstLineIndent=0,
            textColor=colors.HexColor(f"#{INK}"),
        ),
    ]
    return [Paragraph("目录", styles["toc_title"]), toc, PageBreak()]


def paragraph_cell(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(inline_markup(text), style)


def pdf_table(block: Block, styles: dict[str, ParagraphStyle]) -> Table:
    rows = block.rows
    columns = max(len(row) for row in rows)
    width = A4[0] - 4.7 * cm
    column_widths = [width / columns] * columns
    data = []
    for row in rows:
        padded = tuple(row) + ("",) * (columns - len(row))
        data.append([paragraph_cell(value, styles["body"]) for value in padded])
    table = Table(data, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{LIGHT_BLUE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{MID_GRAY}")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def blocks_to_pdf_story(
    blocks: list[Block],
    styles: dict[str, ParagraphStyle],
    *,
    asset_base: Path,
) -> list:
    story: list = []
    list_kind: str | None = None
    list_counter = 0
    for block in blocks:
        if block.kind not in {"bullet", "number"}:
            list_kind = None
            list_counter = 0
        if block.kind == "rule":
            continue
        if block.kind == "heading":
            paragraph = Paragraph(inline_markup(block.text), styles[f"h{block.level}"])
            appendix_detail = block.level == 2 and re.match(r"^[BC]\.\d+", block.text)
            if block.text != "文档识别信息" and not appendix_detail:
                paragraph._toc_level = min(block.level - 1, 1)
            story.append(paragraph)
            continue
        if block.kind == "paragraph":
            story.append(Paragraph(inline_markup(block.text), styles["body"]))
            continue
        if block.kind in {"bullet", "number"}:
            if list_kind != block.kind:
                list_kind = block.kind
                list_counter = 0
            list_counter += 1
            bullet = "•" if block.kind == "bullet" else f"{list_counter}."
            style = ParagraphStyle(
                f"CopyrightList{block.kind}",
                parent=styles["body"],
                leftIndent=18,
                firstLineIndent=-12,
                spaceAfter=3,
            )
            story.append(Paragraph(inline_markup(block.text), style, bulletText=bullet))
            continue
        if block.kind == "callout":
            callout = Table(
                [[Paragraph(inline_markup(block.text), styles["body"])]],
                colWidths=[A4[0] - 4.7 * cm],
            )
            callout.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF5FA")),
                        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor(f"#{BLUE}")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([callout, Spacer(1, 5)])
            continue
        if block.kind == "table":
            story.extend([pdf_table(block, styles), Spacer(1, 6)])
            continue
        if block.kind == "code":
            story.append(XPreformatted(html.escape(block.text or " "), styles["code"]))
            continue
        if block.kind == "image":
            image_path = (asset_base / block.language).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(f"manual image not found: {image_path}")
            image = ReportLabImage(str(image_path))
            image._restrictSize(A4[0] - 4.7 * cm, 17.0 * cm)
            story.append(image)
            if block.text:
                caption_style = ParagraphStyle(
                    "CopyrightImageCaption",
                    parent=styles["body"],
                    alignment=TA_CENTER,
                    fontSize=8,
                    textColor=colors.HexColor(f"#{MUTED}"),
                    spaceAfter=7,
                )
                story.append(Paragraph(inline_markup(block.text), caption_style))
            continue
        raise ValueError(f"unsupported PDF block: {block.kind}")
    return story


def build_pdf(
    source: str,
    output: Path,
    *,
    rights_holder: str,
    source_dir: Path,
) -> None:
    register_pdf_font()
    styles = pdf_styles()
    blocks = parse_markdown(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = CopyrightPdfTemplate(output, styles)
    story = (
        cover_story(styles, rights_holder=rights_holder)
        + toc_story(styles)
        + blocks_to_pdf_story(blocks, styles, asset_base=source_dir)
    )
    document.multiBuild(story)


def document_text(doc: Document) -> str:
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def check_outputs(source: str, docx_path: Path, pdf_path: Path) -> dict[str, int]:
    if not docx_path.is_file() or docx_path.stat().st_size < 20_000:
        raise ValueError(f"DOCX missing or unexpectedly small: {docx_path}")
    if not pdf_path.is_file() or pdf_path.stat().st_size < 20_000:
        raise ValueError(f"PDF missing or unexpectedly small: {pdf_path}")
    if pdf_path.read_bytes()[:4] != b"%PDF":
        raise ValueError("PDF signature is invalid")
    doc = Document(docx_path)
    visible = document_text(doc)
    for heading in REQUIRED_HEADINGS:
        if heading not in visible:
            raise ValueError(f"DOCX missing required heading: {heading}")
    if doc.core_properties.title != MANUAL_TITLE:
        raise ValueError("DOCX title metadata is inconsistent")
    blocks = parse_markdown(source)
    stats = {
        "source_lines": len(source.splitlines()),
        "source_characters": len(source),
        "headings": sum(block.kind == "heading" for block in blocks),
        "tables": sum(block.kind == "table" for block in blocks),
        "code_blocks": sum(block.kind == "code" for block in blocks),
        "images": sum(block.kind == "image" for block in blocks),
        "docx_paragraphs": len(doc.paragraphs),
        "docx_tables": len(doc.tables),
        "docx_bytes": docx_path.stat().st_size,
        "pdf_bytes": pdf_path.stat().st_size,
        "rights_holder_placeholders": source.count(SUBMISSION_PLACEHOLDER),
    }
    if stats["headings"] < 60:
        raise ValueError("registration manual has too few headings")
    if stats["tables"] < 15:
        raise ValueError("registration manual has too few tables")
    if stats["code_blocks"] < 25:
        raise ValueError("registration manual has too few command examples")
    if stats["images"] < 15 or len(doc.inline_shapes) < 15:
        raise ValueError("registration manual has too few operation/design images")
    return stats


def main() -> int:
    args = parse_args()
    source = args.input.read_text(encoding="utf-8")
    validate_source(source)
    rights_holder = RIGHTS_HOLDER
    if args.rights_holder is not None:
        rights_holder = args.rights_holder.strip()
        if not rights_holder or "\n" in rights_holder or "\r" in rights_holder:
            raise ValueError("rights holder must be a non-empty single-line value")
        if rights_holder == SUBMISSION_PLACEHOLDER:
            raise ValueError("rights holder must replace the submission placeholder")
    if SUBMISSION_PLACEHOLDER in source:
        source = source.replace(SUBMISSION_PLACEHOLDER, rights_holder, 1)
    elif rights_holder != RIGHTS_HOLDER:
        source = source.replace(RIGHTS_HOLDER, rights_holder, 1)
    validate_source(source)
    source_dir = args.input.resolve().parent
    build_docx(source, args.docx, rights_holder=rights_holder, source_dir=source_dir)
    build_pdf(source, args.pdf, rights_holder=rights_holder, source_dir=source_dir)
    stats = check_outputs(source, args.docx, args.pdf)
    print(f"built_docx: {args.docx}")
    print(f"built_pdf: {args.pdf}")
    if args.check:
        for key, value in stats.items():
            print(f"{key}: {value}")
        print("copyright manual structural checks: PASS")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # pragma: no cover - CLI failure boundary
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc
